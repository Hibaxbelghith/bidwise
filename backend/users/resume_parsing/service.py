from __future__ import annotations

import logging
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Iterable

from .cleaners import MAX_PARSED_TEXT_CHARS, clean_resume_text_result
from .exceptions import (
    EmptyResumeTextError,
    EncryptedResumeError,
    ResumeParserDependencyMissing,
    ResumeParsingError,
    UnsupportedResumeFormat,
)


logger = logging.getLogger(__name__)

RAW_EXTRACTION_CHAR_LIMIT = MAX_PARSED_TEXT_CHARS * 2
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


@dataclass(frozen=True)
class ParsedResume:
    text: str
    parser: str
    truncated: bool


def _filename_for(file_obj, filename=None) -> str:
    return str(filename or getattr(file_obj, "name", "") or "")


def _extension_for(file_obj, filename=None) -> str:
    return Path(_filename_for(file_obj, filename)).suffix.lower()


@contextmanager
def _opened_binary(file_obj) -> Iterable[BinaryIO]:
    close_after = False
    original_position = None

    if hasattr(file_obj, "open"):
        file_obj.open("rb")
        close_after = True
        stream = file_obj
    else:
        stream = file_obj

    try:
        try:
            original_position = stream.tell()
        except Exception:
            original_position = None
        try:
            stream.seek(0)
        except Exception:
            pass
        yield stream
    finally:
        try:
            stream.seek(original_position or 0)
        except Exception:
            pass
        if close_after:
            try:
                file_obj.close()
            except Exception:
                pass


def _append_bounded(parts: list[str], text: str) -> bool:
    if text:
        parts.append(text)
    return sum(len(part) for part in parts) >= RAW_EXTRACTION_CHAR_LIMIT


def _extract_pdf_text(stream) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise ResumeParserDependencyMissing("pypdf is required for PDF resume parsing.") from exc

    try:
        reader = PdfReader(stream, strict=False)
    except Exception as exc:
        raise ResumeParsingError("PDF could not be opened.") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            decrypted = reader.decrypt("")
        except Exception as exc:
            raise EncryptedResumeError("Encrypted PDF resumes are not supported.") from exc
        if not decrypted:
            raise EncryptedResumeError("Encrypted PDF resumes are not supported.")

    parts: list[str] = []
    try:
        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""
            except PdfReadError as exc:
                raise ResumeParsingError("PDF text extraction failed.") from exc
            if _append_bounded(parts, page_text):
                break
    except EncryptedResumeError:
        raise
    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError("PDF text extraction failed.") from exc

    return "\n".join(parts)


def _extract_docx_text(stream) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeParserDependencyMissing("python-docx is required for DOCX resume parsing.") from exc

    try:
        document = Document(stream)
    except Exception as exc:
        raise ResumeParsingError("DOCX could not be opened.") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        if _append_bounded(parts, paragraph.text):
            return "\n".join(parts)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _append_bounded(parts, paragraph.text):
                        return "\n".join(parts)

    return "\n".join(parts)


def parse_resume_file(
    file_obj,
    *,
    filename: str | None = None,
    max_length: int = MAX_PARSED_TEXT_CHARS,
) -> ParsedResume:
    if not file_obj:
        raise UnsupportedResumeFormat("Resume file is missing.")

    extension = _extension_for(file_obj, filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedResumeFormat("Only PDF and DOCX resumes are supported.")

    with _opened_binary(file_obj) as stream:
        if extension == ".pdf":
            raw_text = _extract_pdf_text(stream)
            parser = "pypdf"
        elif extension == ".docx":
            raw_text = _extract_docx_text(stream)
            parser = "python-docx"
        else:  # pragma: no cover - guarded above
            raise UnsupportedResumeFormat("Only PDF and DOCX resumes are supported.")

    cleaned = clean_resume_text_result(raw_text, max_length=max_length)
    if not cleaned.text:
        raise EmptyResumeTextError("No extractable text found in resume.")

    logger.info(
        "resume parsing completed parser=%s chars=%s truncated=%s",
        parser,
        len(cleaned.text),
        cleaned.truncated,
    )
    return ParsedResume(text=cleaned.text, parser=parser, truncated=cleaned.truncated)
