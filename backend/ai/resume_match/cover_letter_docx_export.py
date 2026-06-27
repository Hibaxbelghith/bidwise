from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


COVER_LETTER_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
COVER_LETTER_FONT = "Arial"


def generate_cover_letter_docx(
    *,
    evidence: dict[str, Any],
    cover_letter_markdown: str,
    verified_email: str = "",
) -> bytes:
    """Build a clean editable DOCX from an already generated cover letter."""
    letter_text = _extract_main_cover_letter(cover_letter_markdown)
    if not letter_text:
        raise ValueError("Cover letter content is empty.")

    doc = Document()
    _configure_document(doc)

    paragraphs = _normalize_candidate_header(
        _split_letter_paragraphs(letter_text),
        evidence=evidence,
        verified_email=verified_email,
    )
    for index, paragraph_text in enumerate(paragraphs):
        _add_letter_paragraph(doc, paragraph_text, index=index)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = COVER_LETTER_FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), COVER_LETTER_FONT)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(8)


def _set_run_font(run) -> None:
    run.font.name = COVER_LETTER_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), COVER_LETTER_FONT)
    run.font.size = Pt(10.5)


def _add_letter_paragraph(doc: Document, text: str, *, index: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.space_after = _paragraph_space_after(text, index)
    run = paragraph.add_run(text)
    _set_run_font(run)
    if index == 0:
        run.bold = True


def _paragraph_space_after(text: str, index: int) -> Pt:
    normalized = _normalize_heading(text)
    if index == 0:
        return Pt(2)
    if normalized in {"bonjour", "dear hiring team", "dear hiring manager"}:
        return Pt(12)
    if normalized.startswith("objet") or normalized.startswith("subject"):
        return Pt(12)
    if normalized in {"cordialement", "sincerely", "best regards"}:
        return Pt(4)
    if len(text) < 90 and (
        "@" in text
        or "[date]" in text.lower()
        or "equipe recrutement" in normalized
        or "hiring team" in normalized
    ):
        return Pt(4)
    return Pt(8)


def _extract_main_cover_letter(markdown: str) -> str:
    text = str(markdown or "").replace("\r\n", "\n").strip()
    if not text:
        return ""

    lines = []
    in_letter = False
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        heading = re.match(r"^#{1,3}\s+(.+?)\s*$", line)
        normalized_heading = _normalize_heading(heading.group(1) if heading else line)

        if normalized_heading in {"lettre de motivation", "cover letter"}:
            in_letter = True
            continue
        if normalized_heading in {
            "version courte pour candidature en ligne",
            "short version for online applications",
            "notes de personnalisation",
            "personalization notes",
        }:
            break
        if not in_letter and heading:
            continue
        if in_letter or not heading:
            cleaned = _strip_markdown(raw_line)
            lines.append(cleaned)

    result = "\n".join(lines).strip()
    # If the model did not include headings, keep the full text but remove known trailing sections.
    if not result:
        result = re.split(
            r"(?im)^#{1,3}\s*(?:version courte pour candidature en ligne|short version for online applications|notes de personnalisation|personalization notes)\s*$",
            text,
            maxsplit=1,
        )[0].strip()
        result = re.sub(r"(?im)^#{1,3}\s*(?:lettre de motivation|cover letter)\s*$", "", result).strip()
    return result


def _split_letter_paragraphs(text: str) -> list[str]:
    normalized_text = _insert_missing_header_breaks(str(text or "").strip())
    lines = [_strip_markdown(line).strip() for line in normalized_text.split("\n")]
    lines = [line for line in lines if line]

    paragraphs: list[str] = []
    buffer: list[str] = []
    in_body = False
    for line in lines:
        normalized = _normalize_heading(line)
        is_header_line = (
            not in_body
            and (
                len(line) < 95
                or "@" in line
                or "[telephone]" in normalized
                or "[phone]" in normalized
                or "[date]" in normalized
                or normalized.startswith("objet")
                or normalized.startswith("subject")
                or normalized in {"bonjour", "dear hiring team", "dear hiring manager"}
            )
        )

        if normalized in {"bonjour", "dear hiring team", "dear hiring manager"}:
            if buffer:
                paragraphs.append(" ".join(buffer))
                buffer = []
            paragraphs.append(line)
            in_body = True
            continue

        if is_header_line:
            if buffer:
                paragraphs.append(" ".join(buffer))
                buffer = []
            paragraphs.append(line)
            continue

        if normalized in {"cordialement", "sincerely", "best regards"}:
            if buffer:
                paragraphs.append(" ".join(buffer))
                buffer = []
            paragraphs.append(line)
            continue

        if in_body:
            buffer.append(line)
            if line.endswith((".", "!", "?")) and len(" ".join(buffer)) > 180:
                paragraphs.append(" ".join(buffer))
                buffer = []
        else:
            paragraphs.append(line)

    if buffer:
        paragraphs.append(" ".join(buffer))
    return paragraphs


def _normalize_candidate_header(
    paragraphs: list[str],
    *,
    evidence: dict[str, Any],
    verified_email: str = "",
) -> list[str]:
    if not paragraphs:
        return paragraphs

    profile = _dict(evidence.get("profile")) if isinstance(evidence, dict) else {}
    contact = _dict(profile.get("contact"))
    name = _clean_header_value(contact.get("full_name"))
    email = _clean_verified_email(verified_email) or _clean_header_value(contact.get("email"))
    locations = _list_values(profile.get("locations"))
    location = locations[0] if locations else ""

    first_block_has_contact = any("@" in item for item in paragraphs[:3])
    if not (name or email or location or first_block_has_contact):
        return paragraphs

    date_index = next(
        (
            index
            for index, item in enumerate(paragraphs[:6])
            if _looks_like_date_placeholder_or_date(item)
        ),
        None,
    )
    recipient_index = next(
        (
            index
            for index, item in enumerate(paragraphs[:6])
            if _normalize_heading(item) in {"equipe recrutement", "hiring team"}
        ),
        None,
    )

    replace_until = 1
    if date_index is not None:
        replace_until = date_index
    elif recipient_index is not None:
        replace_until = recipient_index

    fallback_name = _extract_name_before_email(paragraphs[0])
    header = [
        name or fallback_name or paragraphs[0],
    ]
    if location:
        header.append(location)
    contact_parts = []
    if email:
        contact_parts.append(email)
    elif "@" in paragraphs[0]:
        extracted_email = _extract_email(paragraphs[0])
        if extracted_email:
            contact_parts.append(extracted_email)
    contact_parts.append("[Téléphone]")
    if contact_parts:
        header.append(" | ".join(contact_parts))

    return header + paragraphs[replace_until:]


def _insert_missing_header_breaks(text: str) -> str:
    fixed = str(text or "")
    fixed = re.sub(r"(\S)\s+(\[[Dd]ate\])", r"\1\n\2", fixed)
    fixed = re.sub(r"(\[[Dd]ate\])\s+(Équipe recrutement|Equipe recrutement|Hiring Team)", r"\1\n\n\2", fixed)
    fixed = re.sub(r"(\[[Ll]ien Portfolio ou GitHub\])\s+(\[[Dd]ate\])", r"\1\n\2", fixed)
    fixed = re.sub(r"(\[[Pp]ortfolio or GitHub link\])\s+(\[[Dd]ate\])", r"\1\n\2", fixed)
    fixed = re.sub(r"(\[[Tt]éléphone\]|\[[Tt]elephone\]|\[[Pp]hone\])\s+(\[Lien Portfolio ou GitHub\]|\[Portfolio or GitHub link\])", r"\1 | \2", fixed)
    fixed = re.sub(r"(\| \[[^]]+\])\s+(\[[Dd]ate\])", r"\1\n\2", fixed)
    fixed = re.sub(r"(Équipe recrutement|Equipe recrutement|Hiring Team)\s+([A-Z][^\n]{2,80})", r"\1\n\2", fixed)
    fixed = re.sub(r"([A-Z][A-Z '\-]{4,})\s+(Objet\s*:|Subject\s*:)", r"\1\n\2", fixed)
    fixed = re.sub(r"([A-Z][A-Z '\-]{4,})\s+(Bonjour,|Dear )", r"\1\n\n\2", fixed)
    return fixed


def _looks_like_date_placeholder_or_date(value: str) -> bool:
    text = str(value or "").strip().lower()
    if "[date]" in text:
        return True
    return bool(re.search(r"\b(?:janvier|février|fevrier|mars|avril|mai|juin|juillet|août|aout|septembre|octobre|novembre|décembre|decembre|january|february|march|april|june|july|august|september|october|november|december)\b", text))


def _extract_email(value: str) -> str:
    match = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", str(value or ""))
    return match.group(0) if match else ""


def _clean_verified_email(value: Any) -> str:
    email = str(value or "").strip()
    return email if re.fullmatch(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", email) else ""


def _extract_name_before_email(value: str) -> str:
    text = str(value or "")
    email = _extract_email(text)
    if not email:
        return ""
    prefix = text.split(email, 1)[0].strip()
    # In fused LLM headers, the prefix may contain "Name City, Country".
    if "," in prefix:
        prefix = prefix.split(",", 1)[0].strip()
        words = prefix.split()
        if len(words) > 2:
            prefix = " ".join(words[:2])
    return prefix


def _clean_header_value(value: Any) -> str:
    text = str(value or "").strip()
    return "" if text.lower() in {"not visible", "not visible in resume"} else text


def _dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list_values(value: Any) -> list[str]:
    if isinstance(value, str):
        raw = [value]
    elif isinstance(value, (list, tuple, set)):
        raw = list(value)
    else:
        raw = []
    return [str(item).strip() for item in raw if str(item or "").strip()]


def _strip_markdown(value: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^\s*[-*]\s+", "", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    return text.strip()


def _normalize_heading(value: str) -> str:
    import unicodedata

    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text)
