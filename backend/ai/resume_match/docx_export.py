from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from django.conf import settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_TEMPLATE_PATH = Path(settings.BASE_DIR) / "assets" / "templates" / "wall_street_oasis_resume.docx"
RESUME_FONT = "Times New Roman"
RESUME_TEXT_SIZE = Pt(10)
RESUME_SECTION_SIZE = Pt(11)
RESUME_NAME_SIZE = Pt(16)
RESUME_RIGHT_TAB = Inches(7.25)


def generate_ats_cv_docx(
    *,
    evidence: dict[str, Any],
    optimization_markdown: str,
    ats_document: dict[str, Any] | None = None,
    verified_email: str = "",
) -> bytes:
    """Build an editable ATS CV using the fixed Wall Street Oasis-style template."""
    if isinstance(ats_document, dict) and ats_document:
        return _generate_structured_ats_cv_docx(
            ats_document,
            evidence=evidence,
            verified_email=verified_email,
        )

    # Fallback path: keep the export available even if the dedicated LLM CV
    # generator is unavailable. This uses the prior optimization markdown.
    parsed = _parse_optimization_markdown(optimization_markdown)
    profile = _dict(evidence.get("profile"))
    resume = _dict(evidence.get("resume"))
    opportunity = _dict(evidence.get("opportunity"))

    doc = _new_template_document()
    _clear_document(doc)
    _configure_document(doc)

    contact = _dict(profile.get("contact"))
    name = _text(contact.get("full_name")) or "[Candidate Name]"
    email = _clean_email(verified_email) or _text(contact.get("email")) or "[Email]"
    locations = _list(profile.get("locations"))
    phone = "[Phone]"

    target_title = (
        parsed.get("headline")
        or _first(_list_values(profile.get("target_roles")))
        or _text(resume.get("canonical_role"))
        or _text(opportunity.get("title"))
        or "Target Role"
    )

    _add_name_header(doc, name, email=email, phone=phone, location=_first(locations) or "[City, Country]")
    _add_section(doc, "Target Title")
    _add_paragraph(doc, target_title)

    if parsed.get("summary"):
        _add_section(doc, "Professional Summary")
        _add_paragraph(doc, parsed["summary"])

    skills = parsed.get("skills") or _list_values(resume.get("skills"))[:10]
    tools = _list_values(resume.get("tools"))[:8]
    if skills or tools:
        _add_section(doc, "Core Skills")
        if skills:
            _add_paragraph(doc, ", ".join(skills[:14]))
        if tools:
            _add_paragraph(doc, "Tools: " + ", ".join(tools))

    bullets = parsed.get("experience_bullets") or []
    if bullets:
        _add_section(doc, "Experience Highlights")
        for bullet in bullets[:8]:
            _add_bullet(doc, bullet)

    resume_excerpt = _text(resume.get("summary_excerpt"))
    if resume_excerpt and not bullets:
        _add_section(doc, "Resume Evidence")
        _add_paragraph(doc, resume_excerpt[:900])

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _generate_structured_ats_cv_docx(
    document_data: dict[str, Any],
    *,
    evidence: dict[str, Any],
    verified_email: str = "",
) -> bytes:
    doc = _new_template_document()
    _clear_document(doc)
    _configure_document(doc)

    is_french = _document_looks_french(document_data)
    labels = _section_labels(is_french)
    name = _text(document_data.get("candidate_name")) or "[Candidate Name]"
    contact_line = (
        _contact_line_from_evidence(evidence, is_french, verified_email=verified_email)
        or _replace_email_in_contact_line(_text(document_data.get("contact_line")), verified_email)
        or "[City] | [Email] | [Phone]"
    )
    target_title = _text(document_data.get("target_title"))
    _add_centered_header(doc, name, contact_line, target_title=target_title)

    summary = _text(document_data.get("professional_summary"))
    if summary:
        _add_section(doc, labels["professional_summary"])
        _add_paragraph(doc, summary)

    skills_sections = document_data.get("skills_sections")
    if isinstance(skills_sections, list) and skills_sections:
        _add_section(doc, labels["core_skills"])
        for section in skills_sections[:4]:
            if not isinstance(section, dict):
                continue
            title = _text(section.get("title"))
            items = _list_values(section.get("items"))[:8]
            if title and items:
                _add_paragraph(doc, f"{title}: {', '.join(items)}")

    tools = _list_values(document_data.get("tools"))[:10]
    if tools and not _skills_already_include_tools(skills_sections):
        _add_section(doc, labels["tools"])
        _add_paragraph(doc, ", ".join(tools))

    experience = document_data.get("experience")
    if isinstance(experience, list) and experience:
        _add_section(doc, labels["professional_experience"])
        for entry in experience[:4]:
            if not isinstance(entry, dict):
                continue
            title = _text(entry.get("title"))
            company = _text(entry.get("company"))
            location = _text(entry.get("location"))
            dates = _text(entry.get("dates"))
            if title or company or location or dates:
                _add_experience_heading(
                    doc,
                    title=title,
                    company=company,
                    location=location,
                    dates=dates,
                )
            for bullet in _list_values(entry.get("bullets"))[:6]:
                _add_bullet(doc, bullet)

    education = _list_values(document_data.get("education"))[:4]
    if education:
        _add_section(doc, labels["education"])
        for item in education:
            _add_paragraph(doc, item)

    languages = _list_values(document_data.get("languages"))[:6]
    if languages:
        _add_section(doc, labels["languages"])
        _add_paragraph(doc, ", ".join(languages))

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _section_labels(is_french: bool) -> dict[str, str]:
    if is_french:
        return {
            "target_title": "Titre",
            "professional_summary": "Profil",
            "core_skills": "Compétences",
            "tools": "Outils",
            "professional_experience": "Expérience",
            "education": "Formation",
            "languages": "Langues",
        }
    return {
        "target_title": "Target Title",
        "professional_summary": "Professional Summary",
        "core_skills": "Core Skills",
        "tools": "Tools",
        "professional_experience": "Professional Experience",
        "education": "Education",
        "languages": "Languages",
    }


def _document_looks_french(document_data: dict[str, Any]) -> bool:
    text_parts = [
        document_data.get("target_title"),
        document_data.get("professional_summary"),
        *_list_values(document_data.get("languages")),
    ]
    for section in document_data.get("skills_sections") if isinstance(document_data.get("skills_sections"), list) else []:
        if isinstance(section, dict):
            text_parts.append(section.get("title"))
            text_parts.extend(_list_values(section.get("items")))
    text = " ".join(str(part or "") for part in text_parts).lower()
    return any(marker in text for marker in (" comptable", " compétences", " expérience", " saisie", " français", " déclar"))


def _contact_line_from_evidence(evidence: dict[str, Any], is_french: bool, *, verified_email: str = "") -> str:
    profile = _dict(evidence.get("profile")) if isinstance(evidence, dict) else {}
    contact = _dict(profile.get("contact"))
    email = _clean_email(verified_email) or _text(contact.get("email")) or "[Email]"
    locations = _list_values(profile.get("locations"))
    location = locations[0] if locations else ("[Ville]" if is_french else "[City]")
    phone = "[Téléphone]" if is_french else "[Phone]"
    return f"{location} | {email} | {phone}"

def _clean_email(value: Any) -> str:
    text = str(value or "").strip()
    return text if re.fullmatch(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text) else ""


def _replace_email_in_contact_line(contact_line: str, verified_email: str) -> str:
    email = _clean_email(verified_email)
    text = str(contact_line or "").strip()
    if not text or not email:
        return text
    if re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text):
        return re.sub(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", email, text, count=1)
    return f"{text} | {email}"


def _skills_already_include_tools(skills_sections: Any) -> bool:
    if not isinstance(skills_sections, list):
        return False
    for section in skills_sections:
        if not isinstance(section, dict):
            continue
        title = _text(section.get("title")).lower()
        if any(marker in title for marker in ("outil", "tool", "logiciel", "software")):
            return True
    return False


def _new_template_document() -> Document:
    if DEFAULT_TEMPLATE_PATH.exists():
        return Document(str(DEFAULT_TEMPLATE_PATH))
    return Document()


def _clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    styles = doc.styles
    styles["Normal"].font.name = RESUME_FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), RESUME_FONT)
    styles["Normal"].font.size = RESUME_TEXT_SIZE
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(0)
    styles["Normal"].paragraph_format.line_spacing = 1.0


def _add_name_header(doc: Document, name: str, *, email: str, phone: str, location: str) -> None:
    _add_centered_header(doc, name, f"{location} | {email} | {phone}")


def _add_centered_header(doc: Document, name: str, contact_line: str, *, target_title: str = "") -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(name)
    run.bold = True
    _set_run_font(run, size=RESUME_NAME_SIZE)

    if target_title:
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(0)
        title_run = title.add_run(target_title)
        title_run.bold = True
        _set_run_font(title_run, size=Pt(11))

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5)
    run = contact.add_run(contact_line)
    _set_run_font(run, size=Pt(10))


def _add_section(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(2)
    _add_bottom_border(paragraph)
    run = paragraph.add_run(title.upper())
    run.bold = True
    _set_run_font(run, size=RESUME_SECTION_SIZE)


def _add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(str(text or "").strip())
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        _set_run_font(run)


def _add_role_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(str(text or "").strip())
    run.bold = True
    _set_run_font(run)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)


def _add_experience_heading(
    doc: Document,
    *,
    title: str,
    company: str,
    location: str,
    dates: str,
) -> None:
    first_left = company or title
    first_right = location
    second_left = title if company else ""
    second_right = dates

    if first_left or first_right:
        paragraph = _add_tabbed_paragraph(doc, first_left, first_right)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(0)
        if paragraph.runs:
            paragraph.runs[0].bold = True
        if len(paragraph.runs) >= 3:
            paragraph.runs[-1].italic = True

    if second_left or second_right:
        paragraph = _add_tabbed_paragraph(doc, second_left, second_right)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        if paragraph.runs:
            paragraph.runs[0].italic = True
        if len(paragraph.runs) >= 3:
            paragraph.runs[-1].bold = True


def _add_tabbed_paragraph(doc: Document, left: str, right: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(RESUME_RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    left_run = paragraph.add_run(str(left or "").strip())
    _set_run_font(left_run)
    if right:
        tab_run = paragraph.add_run("\t")
        _set_run_font(tab_run)
        right_run = paragraph.add_run(str(right or "").strip())
        _set_run_font(right_run)
    return paragraph


def _add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("• " + _strip_bullet(text))
    _set_run_font(run)


def _set_run_font(run, *, size=RESUME_TEXT_SIZE) -> None:
    run.font.name = RESUME_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), RESUME_FONT)
    run.font.size = size


def _add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")


def _parse_optimization_markdown(markdown: str) -> dict[str, Any]:
    sections = _markdown_sections(markdown)
    headline = _first_option(
        sections.get("titre de profil proposé")
        or sections.get("updated headline")
        or ""
    )
    summary = _plain_text(
        sections.get("résumé professionnel optimisé")
        or sections.get("updated summary")
        or sections.get("résumé professionnel recommandé")
        or sections.get("recommended professional summary")
        or ""
    )
    skills = _lines(
        sections.get("section compétences à renforcer")
        or sections.get("updated core skills section")
        or sections.get("mots-clés ats à inclure")
        or sections.get("ats keywords to include")
        or ""
    )
    bullets = _lines(
        sections.get("expériences à reformuler pour cette offre")
        or sections.get("resume bullets to strengthen this application")
        or ""
    )
    notes = _lines(
        sections.get("formulations honnêtes pour les compétences partielles")
        or sections.get("honest wording for missing or partial skills")
        or sections.get("notes de nettoyage ats")
        or sections.get("ats cleanup notes")
        or ""
    )
    return {
        "headline": headline,
        "summary": summary,
        "skills": skills,
        "experience_bullets": bullets,
        "notes": notes,
    }


def _markdown_sections(markdown: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {}
    current = ""
    for raw_line in str(markdown or "").replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        heading = re.match(r"^#{1,3}\s+(.+?)\s*$", line)
        if heading:
            current = _normalize_key(heading.group(1))
            sections.setdefault(current, [])
            continue
        if current:
            sections.setdefault(current, []).append(raw_line)
    return {key: "\n".join(value).strip() for key, value in sections.items()}


def _first_option(text: str) -> str:
    for line in _lines(text):
        cleaned = re.sub(r"^option\s+\d+\s*:\s*", "", line, flags=re.IGNORECASE).strip()
        if cleaned:
            return cleaned
    return ""


def _lines(text: str) -> list[str]:
    items: list[str] = []
    for raw_line in str(text or "").splitlines():
        line = _strip_bullet(raw_line)
        if not line or line == "#":
            continue
        if re.match(r"^option\s+\d+\s*:", line, flags=re.IGNORECASE):
            items.append(re.sub(r"^option\s+\d+\s*:\s*", "", line, flags=re.IGNORECASE).strip())
        else:
            items.append(line)
    return [item for item in items if item]


def _plain_text(text: str) -> str:
    return re.sub(r"\s+", " ", "\n".join(_lines(text)) or str(text or "")).strip()


def _strip_bullet(text: str) -> str:
    return re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", str(text or "").strip()).strip()


def _normalize_key(value: str) -> str:
    import unicodedata

    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text)


def _dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _text(value: Any) -> str:
    text = str(value or "").strip()
    return "" if text == "not visible in resume" else text


def _list(value: Any) -> list[str]:
    return _list_values(value)


def _list_values(value: Any) -> list[str]:
    if isinstance(value, str):
        raw = [value]
    elif isinstance(value, (list, tuple, set)):
        raw = list(value)
    else:
        raw = []
    return [str(item).strip() for item in raw if str(item or "").strip()]


def _first(values: list[str]) -> str:
    return values[0] if values else ""
